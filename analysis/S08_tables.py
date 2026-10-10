"""Main tables (Markdown for the manuscript) and supplementary tables (Supplementary_Information.docx with tables and
supplementary figures; Supplementary_Data.xlsx with every table, incl. long ones).
All display-item text 10 pt; supplementary tables left-aligned; italic P (the user's standard).
"""
import re

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font

from common import CACHE, FIG, NEGATIVE, OUT_LABEL, PROJ, RES, follow_up_end

import os
MS = PROJ / os.environ.get("MS_DIR", "manuscript")   # manuscript (AJG version) or manuscript_ClinicalNutrition
SRC = MS / "src"
SRC.mkdir(parents=True, exist_ok=True)
SPECTRUM = ["Gallstones", "Cholecystitis", "AP", "AP_biliary", "AP_nonbiliary", "CP", "Pancreatitis_any"]
SECOND = ["AP_alcohol", "Pancreatic_cancer"]
CONTROLS = ["MASLD", "Cirrhosis", "Dental_caries"] + NEGATIVE
ALTDEF = ["GS_hosp", "GS_or_chole", "Cholecystectomy", "AP_hosp", "CP_fo"]
LAB = {**OUT_LABEL, "AP": "Acute pancreatitis", "AP_biliary": "Gallstone-related acute pancreatitis",
       "AP_nonbiliary": "Non-gallstone acute pancreatitis", "Pancreatitis_any": "Any pancreatitis"}
SIG = {"sig_SSB": "SSB signature", "sig_ASB": "ASB signature", "rsig_SSB": "Residual SSB signature",
       "rsig_ASB": "Residual ASB signature"}
ABBR = ("AP, acute pancreatitis; ASB, artificially sweetened beverages; BMI, body mass index; CI, confidence interval; "
        "CRP, C-reactive protein; FDR, false discovery rate; GGT, gamma-glutamyltransferase; HbA1c, glycated haemoglobin; "
        "HR, hazard ratio; MASLD, metabolic dysfunction-associated steatotic liver disease; SD, standard deviation; "
        "SSB, sugar-sweetened beverages; TG, triglycerides.")


def ci(h, l, u):
    return "—" if pd.isna(h) else f"{h:.2f} ({l:.2f}–{u:.2f})"


def fp(p):
    if pd.isna(p):
        return "—"
    return "<0.001" if p < 0.001 else (f"{p:.3f}" if p < 0.1 else f"{p:.2f}")


def msd(x, dp=1):
    return f"{x.mean():.{dp}f} ({x.std():.{dp}f})"


def miqr(x, dp=1):
    q = x.quantile([0.5, 0.25, 0.75])
    return f"{q.iloc[0]:.{dp}f} ({q.iloc[1]:.{dp}f}–{q.iloc[2]:.{dp}f})"


def npct(m, base=None):
    m = m.fillna(False)
    return f"{int(m.sum()):,} ({100 * m.mean():.1f})"


def md_table(d):
    lines = ["| " + " | ".join(d.columns) + " |", "|" + "---|" * d.shape[1]]
    lines += ["| " + " | ".join("" if (not isinstance(v, str) and pd.isna(v)) else str(v) for v in r) + " |"
              for r in d.itertuples(index=False)]
    return "\n".join(lines)


# ---------------------------------------------------------------- samples for Table 1
d = pd.read_pickle(CACHE / "analysis.pkl")
A = d[d.n_valid_recalls >= 1].copy()
A["start"] = A["first_recall_date"].fillna(A["baseline_date"])
A = A[follow_up_end(A) > A["start"]]
A = A[~(A["Any_cancer"].notna() & (A["Any_cancer"] <= A["start"]))]
sc = pd.read_pickle(CACHE / "signatures.pkl")
B = d.merge(sc, on="eid")
B = B[(B.consortium == 0) & ~B["train"]]
B = B[follow_up_end(B) > B["baseline_date"]]
B = B[~(B["Any_cancer"].notna() & (B["Any_cancer"] <= B["baseline_date"]))]
groups = {"All participants": A, "SSB non-consumers": A[A.SSB == 0], "SSB ≥1 serving/day": A[A.SSB >= 1],
          "ASB non-consumers": A[A.ASB == 0], "ASB ≥1 serving/day": A[A.ASB >= 1], "Independent signature sample": B}


def t1_rows(g, diet=True):
    r = {"Participants, n": f"{len(g):,}", "Age, years": msd(g.age), "Women": npct(g.sex == 0),
         "White ethnicity": npct(g.white == 1), "Townsend deprivation index": msd(g.townsend, 2),
         "College or university degree": npct(g.degree == 1), "Current smoker": npct(g.smoking == 2),
         "Alcohol daily or almost daily": npct(g.alcohol_freq == 1),
         "Physical activity, MET-min/week": miqr(g.met_min, 0), "BMI, kg/m²": msd(g.bmi),
         "Diabetes": npct(g.diabetes_sr == 1), "Triglycerides, mmol/L": miqr(g.tg, 2),
         "HbA1c, mmol/mol": msd(g.hba1c), "C-reactive protein, mg/L": miqr(g.crp, 2)}
    if diet:
        r.update({"Energy intake, kcal/day": msd(g.energy_kj / 4.184, 0), "Recalled alcohol, g/day": miqr(g.alcohol_g),
                  "Total sugars, g/day": msd(g.sugars_g), "SSB, servings/day": msd(g.SSB, 2),
                  "ASB, servings/day": msd(g.ASB, 2), "Fruit juice, servings/day": msd(g.juice, 2),
                  "Tea, cups/day": msd(g.tea), "Coffee, cups/day": msd(g.coffee), "hPDI, score": msd(g.hPDI)})
    else:
        r.update({k: "—" for k in ["Energy intake, kcal/day", "Recalled alcohol, g/day", "Total sugars, g/day",
                                    "SSB, servings/day", "ASB, servings/day", "Fruit juice, servings/day",
                                    "Tea, cups/day", "Coffee, cups/day", "hPDI, score"]})
    return r


t1 = pd.DataFrame({k: t1_rows(g, diet=(k != "Independent signature sample")) for k, g in groups.items()})
t1 = t1.reset_index().rename(columns={"index": "Characteristic"})

# ---------------------------------------------------------------- Table 2 (intake, M2)
im = pd.read_csv(RES / "intake_main.csv")
cat = pd.read_csv(RES / "intake_categories.csv")
py = pd.read_csv(RES / "intake_person_years.csv")
rows = []
for o in SPECTRUM:
    m2 = im[(im.outcome == o) & (im.analysis == "M2")].iloc[0]
    for i, b in enumerate(["SSB", "ASB"]):
        c = cat[(cat.outcome == o) & (cat.exposure == b)].set_index("category")
        rows.append({"Outcome": LAB[o] if i == 0 else "", "Cases": f"{int(m2.events):,}" if i == 0 else "",
                     "Beverage": b, "Non-consumers": "1.00 (ref.)",
                     ">0 to <1 serving/day": ci(c.at["lt1", "hr"], c.at["lt1", "lci"], c.at["lt1", "uci"]),
                     "≥1 serving/day": ci(c.at["ge1", "hr"], c.at["ge1", "lci"], c.at["ge1", "uci"]),
                     "P for trend": fp(c["p_trend"].iloc[0]),
                     "Per 1 serving/day": ci(m2[f"{b}_hr"], m2[f"{b}_lci"], m2[f"{b}_uci"])})
t2 = pd.DataFrame(rows)

# ---------------------------------------------------------------- Table 3 (signatures)
sm = pd.read_csv(RES / "signature_main.csv")
rows = []
for o in SPECTRUM + ["MASLD", "Cirrhosis"]:
    g = sm[sm.outcome == o]
    get = lambda s, mdl, an: g[(g.signature == s) & (g.model == mdl) & (g.analysis == an)].iloc[0]
    r0 = get("sig_SSB", "M2", "per SD")
    rows.append({"Outcome": LAB[o], "Cases/N": f"{int(r0.events):,}/{int(r0.n):,}",
                 **{f"{s.split('_')[1]} signature, {lab}": ci(*get(s, "M2", an)[["hr", "lci", "uci"]])
                    for s in ["sig_SSB", "sig_ASB"]
                    for lab, an in [("Model 2", "per SD"), ("mutually adjusted", "per SD, mutually adjusted")]}})
t3 = pd.DataFrame(rows)

# ---------------------------------------------------------------- write main tables (Markdown)
n_int, n_sig = len(A), len(B)
md = f"""# 表格

**Table 1. Baseline characteristics of participants by intake of sugar-sweetened and artificially sweetened beverages**

{md_table(t1)}

Values are mean (SD), median (IQR) or n (%). The 24-h recall cohort comprises {n_int:,} participants with ≥1 valid recall and no cancer before the first recall; the independent signature sample comprises {n_sig:,} randomly selected Olink participants outside the signature training set (dietary data not required). Beverage intakes are averages over all valid recalls. Abbreviations: hPDI, healthful plant-based diet index; MET, metabolic equivalent of task; {ABBR}

\\pagebreak

**Table 2. Intake of sugar-sweetened and artificially sweetened beverages and risk of biliary and pancreatic diseases**

{md_table(t2)}

Values are HR (95% CI) from Cox models with SSB and ASB in the same model (Model 2): adjusted for age, sex, ethnicity, Townsend deprivation index, education, smoking, alcohol frequency, recalled alcohol intake, physical activity, BMI, diabetes, energy intake, number of recalls and assessment centre. P for trend was obtained by assigning the median intake to each category. Gallstone-related acute pancreatitis was defined as a gallstone record before or within 90 days after the pancreatitis diagnosis, or ICD-10 K85.1. Abbreviations: {ABBR}

\\pagebreak

**Table 3. Proteomic signatures of sugar-sweetened and artificially sweetened beverages and risk of biliary and pancreatic diseases in the independent sample**

{md_table(t3)}

Values are HR (95% CI) per SD of the signature. Model 2 was adjusted for age, sex, ethnicity, Townsend deprivation index, education, smoking, alcohol frequency, physical activity, BMI, diabetes, fasting time and assessment centre (Model 1), plus triglycerides, HbA1c, CRP, GGT, cystatin C and creatinine. Mutually adjusted estimates come from Model 2 including both signatures. MASLD and cirrhosis, the outcomes studied by Zhao et al., are shown for comparison. Abbreviations: {ABBR}
"""
open(SRC / "表格.md", "w", encoding="utf-8").write(md)

# ---------------------------------------------------------------- supplementary tables
SI, DATA = [], []
defs = pd.DataFrame([
    ("Gallstone disease (primary)", "First Occurrence K80 (primary care, hospital, self-report, death)"),
    ("Gallstone disease, hospital", "Hospital inpatient ICD-10 K80"),
    ("Gallstone disease or cholecystectomy", "First Occurrence K80 or OPCS-4 J18"),
    ("Cholecystectomy", "OPCS-4 J18"),
    ("Cholecystitis", "First Occurrence K81"),
    ("Acute pancreatitis (primary)", "First Occurrence K85"),
    ("Acute pancreatitis, hospital", "Hospital inpatient ICD-10 K85"),
    ("Gallstone-related acute pancreatitis", "Acute pancreatitis with a gallstone record (First Occurrence K80) before or within 90 days after the pancreatitis date, or hospital ICD-10 K85.1"),
    ("Non-gallstone acute pancreatitis", "All other acute pancreatitis"),
    ("Alcohol-coded acute pancreatitis", "Hospital ICD-10 K85.2"),
    ("Chronic pancreatitis", "Hospital ICD-10 K86.0, K86.1"),
    ("Other pancreatic disease", "First Occurrence K86"),
    ("Pancreatic cancer", "Cancer registry ICD-10 C25"),
    ("MASLD (positive control)", "First Occurrence K76 or hospital ICD-10 K75.8"),
    ("Liver fibrosis/cirrhosis (positive control)", "First Occurrence K74 or hospital ICD-10 K70.3, K71.7"),
    ("Dental caries (positive control)", "First Occurrence K02"),
    ("Negative controls", "Cataract (H25, H26), inguinal hernia (K40), hearing loss (H90, H91), upper-limb mononeuropathy (G56), varicose veins (I83); First Occurrence"),
    ("Prevalent cases", "Record dated on or before time zero, or self-reported at baseline (gallstones 1162, pancreatitis 1164, liver disease 1158/1604, cataract 1278); prior cholecystectomy excluded for gallbladder outcomes"),
], columns=["Outcome", "Definition"])
SI.append(("Supplementary Table 1. Outcome definitions", defs, "Follow-up ended at the first of outcome, death, loss to follow-up or 31 December 2022."))

fi = pd.read_csv(RES / "flow_intake.csv").set_axis(["Step", "n"], axis=1)
fsg = pd.read_csv(RES / "flow_signature.csv").set_axis(["Step", "n"], axis=1)
flow = pd.concat([fi.assign(Sample="24-h recall cohort"), fsg.assign(Sample="Signature sample")])[["Sample", "Step", "n"]]
flow["n"] = flow["n"].map("{:,}".format)
SI.append(("Supplementary Table 2. Selection of participants", flow, "The beverage–proteome (training/testing) sample comprised 17,183 randomly selected Olink participants with ≥1 valid recall."))

pyt = py[py.outcome.isin(SPECTRUM)].copy()
pyt["Outcome"] = pyt.outcome.map(LAB)
pyt["Category"] = pyt.category.replace({"0": "Non-consumers", ">0 to <1": ">0 to <1 serving/day", ">=1": "≥1 serving/day"})
pyt["Cases/person-years"] = pyt.apply(lambda r: f"{int(r.events):,}/{r.person_years:,.0f}", axis=1)
pyt["Rate per 10,000 person-years"] = (1e4 * pyt.events / pyt.person_years).map("{:.1f}".format)
SI.append(("Supplementary Table 3. Cases and person-years by beverage intake category",
           pyt[["Outcome", "exposure", "Category", "Cases/person-years", "Rate per 10,000 person-years"]].rename(columns={"exposure": "Beverage"}), ""))

rows = []
for o in SPECTRUM + SECOND + ALTDEF + CONTROLS:
    for mdl in ["M1", "M2", "M3"]:
        r = im[(im.outcome == o) & (im.analysis == mdl)]
        if len(r):
            r = r.iloc[0]
            rows.append({"Outcome": LAB[o], "Model": mdl.replace("M", "Model "), "Cases": f"{int(r.events):,}",
                         "SSB per serving/day": ci(r.SSB_hr, r.SSB_lci, r.SSB_uci), "SSB P": fp(r.SSB_p),
                         "ASB per serving/day": ci(r.ASB_hr, r.ASB_lci, r.ASB_uci), "ASB P": fp(r.ASB_p)})
SI.append(("Supplementary Table 4. Beverage intake and all outcomes across adjustment models", pd.DataFrame(rows),
           "Model 1: age, sex, energy intake, number of recalls, centre. Model 2: + ethnicity, Townsend index, education, smoking, alcohol frequency, recalled alcohol, physical activity, BMI, diabetes. Model 3: + TG, HbA1c, CRP, GGT. SSB and ASB in the same model."))

se = pd.read_csv(RES / "intake_sensitivity.csv")
se = se[se.outcome.isin(["Gallstones", "AP", "AP_biliary", "AP_nonbiliary", "CP", "Cholecystitis"])]
se = pd.DataFrame({"Outcome": se.outcome.map(LAB), "Analysis": se.analysis.str.replace(">=", "≥"), "Cases": se.events.map("{:,}".format),
                   "SSB per serving/day": [ci(*x) for x in se[["SSB_hr", "SSB_lci", "SSB_uci"]].to_numpy()],
                   "ASB per serving/day": [ci(*x) for x in se[["ASB_hr", "ASB_lci", "ASB_uci"]].to_numpy()]})
SI.append(("Supplementary Table 5. Sensitivity analyses of beverage intake and biliary and pancreatic diseases", se,
           "All models are Model 2 unless stated. Gallstone-related confounders: hormone replacement therapy, oral contraceptives, parity, cholesterol-lowering and antihypertensive medication, insulin and weight change in the past year."))

fg = pd.read_csv(RES / "finegray.csv")
fg = pd.DataFrame({"Analysis": fg.analysis, "Outcome": fg.outcome.map(LAB), "Exposure": fg.exposure.map(lambda x: SIG.get(x, x)),
                   "Cases": fg.events.map("{:,}".format), "Competing deaths": fg.deaths.map("{:,}".format),
                   "Subdistribution HR (95% CI)": [ci(*x) for x in fg[["shr", "lci", "uci"]].to_numpy()], "P": fg.p.map(fp)})
SI.append(("Supplementary Table 6. Fine–Gray competing-risk models with death as the competing event", fg, "Covariates as in Model 2 of the respective analysis."))

sg = pd.read_csv(RES / "intake_subgroups.csv")
sg = pd.DataFrame({"Outcome": sg.outcome.map(LAB), "Beverage": sg.exposure, "Subgroup": sg.subgroup, "Level": sg.level.str.replace(">=", "≥"),
                   "Cases": sg.events.map("{:,}".format), "HR per serving/day": [ci(*x) for x in sg[["hr", "lci", "uci"]].to_numpy()],
                   "P for interaction": sg.p_interaction.map(fp)})
SI.append(("Supplementary Table 7. Subgroup analyses of beverage intake", sg, "Model 2 fitted within strata; P for interaction from likelihood-ratio tests of beverage × subgroup terms."))

st = pd.read_csv(RES / "intake_substitution.csv")
st = pd.DataFrame({"Outcome": st.outcome.map(LAB), "Replace 1 serving/day of": st["replace"],
                   "With 1 serving/day of": st["with"].replace({"juice": "fruit juice"}),
                   "HR (95% CI)": [ci(*x) for x in st[["hr", "lci", "uci"]].to_numpy()], "P": st.p.map(fp)})
SI.append(("Supplementary Table 8. Substitution analyses", st, "Model 2 including SSB, ASB, water, tea, coffee and fruit juice simultaneously; HRs are differences between beverage coefficients."))

bm = pd.read_csv(RES / "biomarkers.csv")
bm = pd.DataFrame({"Biomarker": bm.biomarker, "Beverage": bm.exposure, "Adjustment": bm.adjustment, "n": bm.n.map("{:,}".format),
                   "% difference per serving/day (95% CI)": [f"{a:.2f} ({b:.2f} to {c:.2f})" for a, b, c in bm[["pct", "lci", "uci"]].to_numpy()],
                   "P": bm.p.map(fp)})
SI.append(("Supplementary Table 9. Beverage intake and baseline blood biomarkers", bm, "Linear regression of log-transformed biomarkers on SSB and ASB (mutually adjusted) with Model 2 covariates and fasting time."))

pa = pd.read_csv(RES / "protein_assoc.csv")
cnt = pa.groupby(["model", "exposure"]).apply(lambda g: pd.Series({"Proteins at FDR <0.05": int((g.fdr < 0.05).sum()),
                                                                   "Positive": int(((g.fdr < 0.05) & (g.beta > 0)).sum())})).reset_index()
cnt.columns = ["Model", "Beverage", "Proteins at FDR <0.05", "Positive"]
SI.append(("Supplementary Table 10. Number of plasma proteins associated with beverage intake", cnt,
           "Main model: Model 2 covariates, fasting time and the other beverage; 'without BMI' removes BMI. Full results are in Supplementary Data."))
rep = pd.read_csv(RES / "natfood_replication.csv")
rep = pd.DataFrame({"Beverage": rep.exposure, "Protein": rep.protein.str.upper(), "β per serving/day": rep.beta.map("{:.3f}".format),
                    "P": rep.p.map(fp), "FDR": rep.fdr.map(fp)})
SI.append(("Supplementary Table 11. Replication of the proteins highlighted by Zhao et al. (Nat Food 2025, Fig. 4)", rep,
           "Proteins as reported in the volcano plots of the original study; estimates from the main model of this study."))
perf = pd.read_csv(RES / "signature_performance.csv")
perf = pd.DataFrame({"Signature": perf.signature.map(SIG), "Target": perf.target, "Training n": perf.n_train.map("{:,}".format),
                     "Testing n": perf.n_test.map("{:,}".format), "Proteins selected": perf.n_proteins,
                     "Pearson r (test)": perf.pearson_r_test.map("{:.3f}".format),
                     "Spearman r with intake (test)": perf.spearman_r_test.map("{:.3f}".format),
                     "r with BMI signature": perf.r_with_BMI_signature.map("{:.2f}".format)})
SI.append(("Supplementary Table 12. Performance of the proteomic signatures", perf,
           "Residual signatures were trained on intake residualised on demographic, lifestyle, adiposity (BMI, waist), liver (ALT, GGT), glycaemic (HbA1c), lipid (TG), kidney (cystatin C, creatinine) and inflammatory (CRP) covariates."))

rows = []
for o in SPECTRUM + SECOND + ALTDEF + CONTROLS:
    g = sm[(sm.outcome == o) & (sm.analysis == "per SD")]
    if not len(g):
        continue
    r = {"Outcome": LAB[o], "Cases/N": f"{int(g.events.iloc[0]):,}/{int(g.n.iloc[0]):,}"}
    for s in ["sig_SSB", "sig_ASB", "rsig_SSB", "rsig_ASB"]:
        for mdl in ["M1", "M2", "M3"]:
            x = g[(g.signature == s) & (g.model == mdl)].iloc[0]
            r[f"{SIG[s]}, {mdl.replace('M', 'Model ')}"] = ci(x.hr, x.lci, x.uci)
    rows.append(r)
s13 = pd.DataFrame(rows)
SI.append(("Supplementary Table 13. Proteomic signatures and all outcomes across adjustment models (HR per SD)",
           s13[["Outcome", "Cases/N"] + [c for c in s13.columns if c.startswith(("SSB", "ASB"))]],
           "Model 1: age, sex, ethnicity, Townsend index, education, smoking, alcohol frequency, physical activity, BMI, diabetes, fasting time, centre. Model 2: + TG, HbA1c, CRP, GGT, cystatin C, creatinine. Model 3: + BMI proteomic signature. Residual signatures are shown in Supplementary Data."))
tt = pd.read_csv(RES / "signature_tertiles.csv")
tt = tt[tt.outcome.isin(SPECTRUM)]
tt = pd.DataFrame({"Outcome": tt.outcome.map(LAB), "Signature": tt.signature.map(SIG), "Model": tt.model.str.replace("M", "Model "),
                   "Tertile": tt.tertile + " vs T1", "HR (95% CI)": [ci(*x) for x in tt[["hr", "lci", "uci"]].to_numpy()],
                   "P for trend": tt.p_trend.map(fp)})
SI.append(("Supplementary Table 14. Signature tertiles and biliary and pancreatic diseases", tt, ""))
ss = pd.read_csv(RES / "signature_sensitivity.csv")
ss = pd.DataFrame({"Outcome": ss.outcome.map(LAB), "Analysis": ss.analysis, "Signature": ss.signature.map(SIG),
                   "Cases": ss.events.map("{:,}".format), "HR per SD (95% CI)": [ci(*x) for x in ss[["hr", "lci", "uci"]].to_numpy()],
                   "P": ss.p.map(fp)})
SI.append(("Supplementary Table 15. Sensitivity analyses of the proteomic signatures", ss,
           "Model 2 unless stated; the gallstone-confounder analysis adds those covariates to Model 1."))
het = pd.read_csv(RES / "heterogeneity_ap.csv")
het = pd.DataFrame({"Analysis": het.analysis, "Exposure": het.exposure.map(lambda x: SIG.get(x, x)),
                    "HR, gallstone-related AP": het.hr_biliary.map("{:.2f}".format),
                    "HR, non-gallstone AP": het.hr_nonbiliary.map("{:.2f}".format),
                    "Ratio of HRs": het.ratio_of_hr.map("{:.2f}".format), "P for heterogeneity": het.p_heterogeneity.map(fp)})
SI.append(("Supplementary Table 16. Heterogeneity between gallstone-related and non-gallstone acute pancreatitis", het,
           "z test on the difference of log HRs, treating the two subtype estimates as independent."))
evt = pd.read_csv(RES / "evalues.csv")
evt = evt[evt.outcome.isin(SPECTRUM)]
evt = pd.DataFrame({"Source": evt.source, "Outcome": evt.outcome.map(LAB), "Exposure": evt.exposure.map(lambda x: SIG.get(x, x)),
                    "HR": evt.hr.map("{:.2f}".format), "E-value (point estimate)": evt.e_value.map("{:.2f}".format),
                    "E-value (CI limit)": evt.e_value_ci.map("{:.2f}".format)})
SI.append(("Supplementary Table 17. E-values for unmeasured confounding", evt, "E-value of 1.00 for the CI limit indicates that the CI includes 1."))

pht = pd.read_csv(RES / "ph_test.csv")
pht = pd.DataFrame({"Analysis": pht.analysis, "Outcome": pht.outcome.map(LAB), "Exposure": pht.term.map(lambda x: SIG.get(x, x)),
                    "P for proportional hazards": pht.p_PH.map(fp)})
tsp = pd.read_csv(RES / "ph_timesplit.csv")
tsp = pd.DataFrame({"Analysis": tsp.analysis, "Outcome": tsp.outcome.map(LAB), "Exposure": tsp.term.map(lambda x: SIG.get(x, x)),
                    "Follow-up period": tsp.period, "Cases": tsp.events, "HR (95% CI)": [ci(*x) for x in tsp[["hr", "lci", "uci"]].to_numpy()]})
SI.append(("Supplementary Table 18. Proportional-hazards tests for the main models", pht,
           "Scaled Schoenfeld residual test with rank-transformed time, exposure terms of Model 2."))
SI.append(("Supplementary Table 19. Time-split hazard ratios for terms with a proportional-hazards P <0.05", tsp,
           "Events in the first 6 years (follow-up censored at 6 years) and after 6 years (participants event-free at 6 years)."))

mrr = pd.read_csv(RES / "cis_mr_results.csv")
mrp = mrr[mrr.method.isin(["Wald ratio", "IVW"])].sort_values("p")
mr_tab = lambda d: pd.DataFrame({"Protein": d.protein, "Outcome (FinnGen R12)": d.outcome, "Method": d.method, "SNPs": d.nsnp,
                                 "OR per SD (95% CI)": [ci(*x) for x in d[["OR", "lci", "uci"]].to_numpy()],
                                 "P": d.p.map(fp), "FDR": d.fdr.map(fp)})
SI.append(("Supplementary Table 20. Cis-pQTL Mendelian randomization of signature proteins (tests with P <0.05)",
           mr_tab(mrp[mrp.p < 0.05]),
           "232 protein–outcome tests (58 proteins with instruments matched in FinnGen); all results are in Supplementary Data 25. "
           "Instruments: conditionally independent cis-pQTLs from the UKB-PPP discovery cohort (-log10 P >=7.3, r2 <0.01)."))
cr = pd.read_csv(RES / "coloc_results.csv")
cr = cr[cr.p12 == 1e-5].assign(Window="Gene ±1 Mb")
cw = pd.read_csv(RES / "coloc_window.csv").assign(Window="Instrument ±100 kb")
cc = pd.concat([cr, cw])
cc = pd.DataFrame({"Protein": cc.protein, "Outcome": cc.outcome, "Window": cc.Window, "SNPs": cc.nsnps,
                   "PP.H3": cc["PP.H3"].map("{:.3f}".format), "PP.H4": cc["PP.H4"].map("{:.3f}".format),
                   "Top shared variant": cc.top_shared_rsid if "top_shared_rsid" in cc else ""})
rmr = pd.read_csv(RES / "cis_mr_region.csv")
SI.append(("Supplementary Table 21. Colocalisation of MENT, CPB1 and IL6R with biliary and pancreatic diseases", cc,
           "coloc.abf with priors p1 = p2 = 1e-4 and p12 = 1e-5 (results with p12 = 5e-6 were similar). PP.H3, distinct causal variants; "
           "PP.H4, shared causal variant. Lead cholelithiasis signals in the regions: CERS2 (MENT), TM4SF4 (CPB1) and ADAR (IL6R)."))

rci = pd.read_csv(RES / "reverse_causation_intake.csv")
a = rci[rci.analysis == "per serving/day"]
rc1 = pd.DataFrame({"Restriction (cumulative)": a.restriction, "Outcome": a.outcome.map(LAB), "Cases": a.events.map("{:,}".format),
                    "SSB per serving/day": [ci(*x) for x in a[["SSB_hr", "SSB_lci", "SSB_uci"]].to_numpy()],
                    "ASB per serving/day": [ci(*x) for x in a[["ASB_hr", "ASB_lci", "ASB_uci"]].to_numpy()]})
b = rci[rci.analysis != "per serving/day"].copy()
b["est"] = [ci(r[f"{r.exposure}_hr"], r[f"{r.exposure}_lci"], r[f"{r.exposure}_uci"]) for _, r in b.iterrows()]
rc1b = pd.DataFrame({"Sample": b.restriction.replace({"All participants (main)": "All with >=2 recalls",
                                                      "+ stable weight in past year": "+ no dietary change, stable weight"}),
                     "Outcome": b.outcome.map(LAB), "Beverage": b.exposure,
                     "Comparison": b.analysis.str.extract(r"\((.*)\)")[0], "Cases": b.events.map("{:,}".format), "HR (95% CI)": b.est})
SI.append(("Supplementary Table 22. Beverage intake and biliary and pancreatic diseases after restrictions addressing reverse causation",
           rc1, "Model 2 with SSB and ASB mutually adjusted. Restrictions are cumulative: no major dietary change in the past 5 years "
           "(field 1538), stable weight in the past year (field 2306), no diabetes, BMI <30 kg/m2."))
SI.append(("Supplementary Table 23. Habitual versus occasional consumption among participants with ≥2 valid recalls", rc1b,
           "Habitual: consumed in every valid recall; occasional: in some recalls; reference: in none. SSB and ASB patterns in the same Model 2."))
rcs = pd.read_csv(RES / "reverse_causation_signature.csv")
rc2 = pd.DataFrame({"Restriction (cumulative)": rcs.restriction, "Outcome": rcs.outcome.map(LAB), "Signature": rcs.signature.map(SIG),
                    "Cases": rcs.events.map("{:,}".format), "HR per SD (95% CI)": [ci(*x) for x in rcs[["hr", "lci", "uci"]].to_numpy()]})
SI.append(("Supplementary Table 24. Proteomic signatures and biliary and pancreatic diseases after restrictions addressing reverse causation",
           rc2, "Model 2 in the independent signature sample; restrictions as in Supplementary Table 22."))
sr = pd.read_csv(RES / "substitution_restricted.csv")
SI[-3] = (SI[-3][0], SI[-3][1], SI[-3][2] + " Substitution estimates in the restricted sample (no dietary change, stable weight, no diabetes) are in Supplementary Data 31.")

# long tables: Excel only
pc = pd.read_csv(RES / "signature_protein_cox.csv")
pcx = pd.DataFrame({"Protein": pc.protein.str.upper(), "Outcome": pc.outcome.map(LAB), "Cases": pc.events,
                    "HR per SD (95% CI)": [ci(*x) for x in pc[["hr", "lci", "uci"]].to_numpy()], "P": pc.p.map(fp),
                    "FDR": pc.fdr.map(fp), "Weight in SSB signature": pc.weight_SSB.round(4), "Weight in ASB signature": pc.weight_ASB.round(4)})
pax = pa.copy()
pax["protein"] = pax.protein.str.upper()
pax = pax.rename(columns={"model": "Model", "exposure": "Beverage", "protein": "Protein", "beta": "β per serving/day", "se": "SE",
                          "p": "P", "fdr": "FDR"})
wts = []
for k in ["sig_SSB", "sig_ASB", "rsig_SSB", "rsig_ASB"]:
    w = pd.read_csv(RES / f"signature_weights_{k}.csv")
    w.columns = ["Protein", "Weight"]
    w.insert(0, "Signature", SIG[k])
    w["Protein"] = w["Protein"].str.upper()
    wts.append(w)
enr = []
for k, lab in [("SSB_FDR", "SSB-related proteins (FDR <0.05)"), ("SSB_signature", "SSB signature proteins"),
               ("ASB_signature", "ASB signature proteins")]:
    f = RES / f"enrichment_{k}.csv"
    e = pd.read_csv(f) if f.exists() and f.stat().st_size > 5 else pd.DataFrame()
    if len(e):
        e.insert(0, "Protein set", lab)
        enr.append(e.rename(columns={"p_adj": "FDR-adjusted P"}))
enr = pd.concat(enr) if enr else pd.DataFrame({"note": ["no enriched terms"]})
full_sig = sm.assign(Outcome=sm.outcome.map(LAB), Signature=sm.signature.map(SIG),
                     **{"HR (95% CI)": [ci(*x) for x in sm[["hr", "lci", "uci"]].to_numpy()]})[
    ["Outcome", "analysis", "model", "Signature", "events", "n", "HR (95% CI)", "p"]].rename(
    columns={"analysis": "Analysis", "model": "Model", "events": "Cases", "n": "N", "p": "P"})

XL_ONLY = [("Supplementary Data 25. Cis-pQTL Mendelian randomization, all estimates", mrr.assign(protein=mrr.protein)),
           ("Supplementary Data 26. Protein-wise associations of SSB and ASB with 2,919 plasma proteins", pax),
           ("Supplementary Data 27. Elastic-net weights of the proteomic signatures", pd.concat(wts)),
           ("Supplementary Data 28. Individual signature proteins and biliary and pancreatic diseases (Model 1)", pcx),
           ("Supplementary Data 29. Pathway enrichment (g:Profiler)", enr),
           ("Supplementary Data 30. All signature–outcome estimates, incl. residual signatures", full_sig),
           ("Supplementary Data 31. Substitution analyses in the restricted sample", sr)]
DATA = [(t, dd) for t, dd, _ in SI] + XL_ONLY

# ---------------------------------------------------------------- Word: supplementary information
P_RE = re.compile(r"(?<![A-Za-z])P(?![A-Za-z])")
doc = Document()
stl = doc.styles["Normal"]; stl.font.name = "Arial"; stl.font.size = Pt(10)


def add_text(par, text, bold=False, italic=False):
    pos = 0
    for m in list(P_RE.finditer(text)) + [None]:
        end = m.start() if m else len(text)
        if end > pos:
            r = par.add_run(text[pos:end]); r.font.size = Pt(10); r.bold = bold or None; r.italic = italic or None
        if m:
            r = par.add_run("P"); r.font.size = Pt(10); r.bold = bold or None; r.italic = True; pos = m.end()


def border(cell, edges):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
    for e in ("top", "bottom", "left", "right"):
        el = OxmlElement(f"w:{e}")
        el.set(qn("w:val"), "single" if e in edges else "nil")
        if e in edges:
            el.set(qn("w:sz"), "8"); el.set(qn("w:color"), "000000")
        b.append(el)
    tcPr.append(b)


def orient(sec, land):
    if land and sec.page_width < sec.page_height:
        sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = sec.page_height, sec.page_width
    if not land and sec.page_width > sec.page_height:
        sec.orientation = WD_ORIENT.PORTRAIT; sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for mg in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, mg, Cm(1.6))


add_text(doc.add_paragraph(), "Supplementary Information", bold=True)
add_text(doc.add_paragraph(), os.environ.get("SI_TITLE", "Sugar-sweetened and artificially sweetened beverages, plasma "
                              "proteomic signatures and risk of biliary and pancreatic diseases: a prospective cohort study in the UK Biobank"))
add_text(doc.add_paragraph(), "Contents: Supplementary Tables 1–24; Supplementary Figures 1–4. Supplementary Data 25–31 "
                              "(full protein-level results) are provided in Supplementary_Data.xlsx.")
first = True
for title, dd, note in SI:
    land = dd.shape[1] > 5 or dd.astype(str).apply(lambda c: c.str.len().max()).sum() > 95
    sec = doc.add_section()
    orient(sec, land)
    add_text(doc.add_paragraph(), title, bold=True)
    t = doc.add_table(rows=len(dd) + 1, cols=dd.shape[1])
    for j, h in enumerate(dd.columns):
        c = t.cell(0, j); c.text = ""; add_text(c.paragraphs[0], str(h), bold=True); border(c, ("top", "bottom"))
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    for i, row in enumerate(dd.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = ""
            add_text(c.paragraphs[0], "" if (not isinstance(v, str) and pd.isna(v)) else str(v))
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            border(c, ("bottom",) if i == len(dd) else ())
    add_text(doc.add_paragraph(), (note + " " if note else "") + "Abbreviations: " + ABBR)

FIGS = [("FigS1_substitution", "Supplementary Figure 1. Substitution analyses. HRs (95% CI) for replacing 1 serving/day of "
         "SSB (a) or ASB (b) with 1 serving/day of another beverage, from Model 2 including all beverages."),
        ("FigS2_subgroups", "Supplementary Figure 2. Subgroup analyses of beverage intake and gallstone disease (a) and acute "
         "pancreatitis (b). HR per serving/day (Model 2); P for interaction from likelihood-ratio tests."),
        ("FigS3_signature_tertiles", "Supplementary Figure 3. Tertiles of the SSB (a) and ASB (b) proteomic signatures and "
         "biliary and pancreatic diseases (Model 2), lowest tertile as reference."),
        ("FigS4_biomarkers", "Supplementary Figure 4. Percentage difference in baseline blood biomarkers per serving/day of SSB "
         "and ASB (Model 2 covariates and fasting time; mutually adjusted).")]
for name, legend in FIGS:
    sec = doc.add_section()
    orient(sec, False)
    doc.add_picture(str(FIG / f"{name}.png"), width=Cm(17))
    add_text(doc.add_paragraph(), legend)
doc.save(MS / "Supplementary_Information.docx")

# ---------------------------------------------------------------- Excel: supplementary data
F10, F10B, LEFT = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True), Alignment(horizontal="left")


def xl(v, bold=False):
    if not isinstance(v, str) or not P_RE.search(v):
        return v
    f = lambda it: InlineFont(rFont="Arial", sz=10, b=bold or None, i=it or None)
    out, pos = [], 0
    for m in P_RE.finditer(v):
        if m.start() > pos:
            out.append(TextBlock(f(False), v[pos:m.start()]))
        out.append(TextBlock(f(True), "P")); pos = m.end()
    if pos < len(v):
        out.append(TextBlock(f(False), v[pos:]))
    return CellRichText(out)


wb = Workbook(); ws0 = wb.active; ws0.title = "Contents"
for i, (title, _) in enumerate(DATA, start=1):
    ws0.cell(row=i, column=1, value=xl(title)).font = F10
ws0.column_dimensions["A"].width = 110
for title, dd in DATA:
    num = re.search(r"(\d+)\.", title).group(1)
    ws = wb.create_sheet(("Suppl. Table " if title.startswith("Supplementary Table") else "Suppl. Data ") + num)
    ws.cell(row=1, column=1, value=xl(title, bold=True)).font = F10B
    for j, c in enumerate(dd.columns, start=1):
        h = ws.cell(row=3, column=j, value=xl(str(c), bold=True)); h.font, h.alignment = F10B, LEFT
    for i, row in enumerate(dd.itertuples(index=False), start=4):
        for j, v in enumerate(row, start=1):
            v = None if (isinstance(v, float) and np.isnan(v)) else (v.item() if hasattr(v, "item") else v)
            c = ws.cell(row=i, column=j, value=v); c.font, c.alignment = F10, LEFT
    for j, c in enumerate(dd.columns, start=1):
        w = max([len(str(c))] + [len(str(v)) for v in dd.iloc[:200, j - 1]]) + 2
        ws.column_dimensions[ws.cell(row=3, column=j).column_letter].width = min(60, max(9, w))
wb.save(MS / "Supplementary_Data.xlsx")
print(t1.to_string(index=False)); print(t2.to_string(index=False)); print(t3.to_string(index=False))
